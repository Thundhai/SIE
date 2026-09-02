"""Real enterprise dataset loader — Real Enterprise Dataset Evaluation
v0.1. Maps a supplied HSE observation/incident workbook export into
SIE's existing enterprise ingestion contract (`RawSafetyEventPayload`),
the exact same type `POST /api/v1/data/ingestion` and every fixture in
`tests/fixtures/` already produces. This module does no ingestion,
validation, or analysis itself — it is purely a mapping layer; every
downstream step (validation, terminology mapping, temporal integrity,
provenance, intelligence, predictive readiness) is the existing,
unmodified SIE pipeline (see `app/validation/real_dataset_evaluation.py`).

**Never commits, embeds, or hardcodes the supplied workbook's own
content.** This module only knows how to read a workbook from a path a
caller supplies at runtime — no sample data, no fixture derived from the
real file, lives in this repository. See
`docs/REAL_ENTERPRISE_DATASET_EVALUATION_GUIDE.md` for the full data-
handling policy.

**Mapping decisions, made once, deliberately, and left auditable here —
never "guessed" at runtime:**

* `event_type` for an Incident-sheet row is the workbook's own
  `Incident Type` value, passed through **unresolved** (not translated
  by this module) — because that column genuinely carries top-level
  domain information the source system itself distinguishes (e.g.
  `NearMiss` is not a subtype of `Injury`; conflating the two by forcing
  every Incident-sheet row to `event_type="INCIDENT"` would misclassify
  every near-miss). Whether a given raw term resolves against SIE's
  existing, **unmodified** `app/intelligence/terminology_mapping.py`
  alias table is entirely the terminology-review stage's job, not this
  loader's -- see that module's own docstring on why guessing is never
  acceptable. A term this workbook uses that the existing table does not
  recognize (e.g. `NearMiss`, `FireIncident`) is a genuine, reportable
  finding, not something this loader silently resolves.
* `event_type` for an Observation-sheet row is set directly to the
  literal string `"OBSERVATION"` — not a guess, but the one piece of
  information the source workbook itself provides unambiguously by
  construction (a dedicated worksheet per record kind, with no competing
  top-level classification column on that sheet the way the Incident
  sheet's own `Incident Type` is). `Category` (the sheet's real safety
  classification -- Unsafe Act, Housekeeping, ...) is passed through
  unresolved as `event_subtype`, for the same reason `Incident Type` is.
* **The Observation sheet's own duplicate `Status` column is never
  resolved for the caller.** Openpyxl exposes two columns both literally
  named `Status`. The first (business open/closed lifecycle) feeds
  `RawSafetyEventPayload.status` directly. The second is preserved
  verbatim in `attributes["status_secondary_raw"]` — **never merged,
  never asserted as authoritative, never treated as a duplicate of the
  first.** (A real finding, not a hypothetical: in the supplied
  workbook the two columns are near-perfectly *inversely* shaped —
  959 "Closed" vs. 959 "REVIEW REQUIRED" — suggesting the second column
  is an internal data-quality review flag from the source system, not a
  restatement of the close/open status. See the evaluation report's own
  "Architecture limitations discovered" section.)
* Every other narrative/free-text column (`Description`, `WitnessStatement`,
  `Root Cause Analysis`, `Corrective Action`, `Preventive Action`,
  `Recommendation`, `Conclusion`, `PeopleInvolved`, `Reported By`,
  `Finding`, `Comment`, `Action Taken`) is preserved on the payload
  (`description` for the sheet's own primary narrative field, everything
  else in `attributes`) — **never dropped, never fabricated when
  missing.** These fields may contain personally identifying
  information (real names appear in this dataset's own incident
  narratives); they are stored exactly like any other enterprise
  ingestion payload always has been (in the application's own database,
  never in source control) and are never read by feature engineering
  (`app/intelligence/features.py` only ever reads structured
  event_type/subtype/status/severity/site/time fields — see
  `tests/test_real_enterprise_dataset_evaluation.py`'s own PII-exclusion
  regression test).
* A row missing a required identifier (`Document No`) or a required
  timestamp (`Date`) still produces a payload (with that field `None`)
  — this loader never silently drops a row. Whether that omission
  results in `REJECTED_INVALID`/`QUARANTINED`/`PARTIAL` is entirely the
  existing, unmodified `app/intelligence/validation.py` pipeline's
  decision, made once the payload actually reaches ingestion.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.intelligence.schemas import RawSafetyEventPayload

#: A single shared source_system: both sheets are two record types
#: within the one supplied workbook/export, not two independent systems.
#: Document-number prefixes (`ALM-HSE-IC-`/`ALM-HSE-OB-`) already keep
#: `(source_system, source_record_id)` identity disjoint between them.
SOURCE_SYSTEM = "alm-hse-xlsx"

_INCIDENT_SHEET_CANDIDATES = ("Incident", "Incidents")
_OBSERVATION_SHEET_CANDIDATES = ("Observation", "Observations")

#: Structural columns this loader depends on existing (by name, after
#: whitespace normalization -- the supplied workbook's own headers carry
#: inconsistent trailing spaces, e.g. `"Date "` on the Incident sheet
#: vs. `"Date"` on the Observation sheet). Any *other* missing/blank
#: cell is a per-row data-quality condition for the existing validation
#: pipeline to classify, not a structural failure this loader reports.
_INCIDENT_REQUIRED_HEADERS = ("Date", "Document No", "Project", "Incident Type")
_OBSERVATION_REQUIRED_HEADERS = ("Document No", "Date", "Project Name", "Category")

_INCIDENT_ATTRIBUTE_COLUMNS = (
    "PeopleInvolved", "WitnessStatement", "ImmediateActionTaken", "Root Cause Analysis",
    "Corrective Action", "Preventive Action", "Recommendation", "Conclusion", "Reported By",
)


@dataclass
class StructureValidationIssue:
    sheet: str
    code: str
    message: str


@dataclass
class RealDatasetLoadResult:
    incident_payloads: list[RawSafetyEventPayload] = field(default_factory=list)
    observation_payloads: list[RawSafetyEventPayload] = field(default_factory=list)
    structure_issues: list[StructureValidationIssue] = field(default_factory=list)
    row_level_skips: list[str] = field(default_factory=list)  # human-readable, no PII -- "<sheet> row <n>: <reason>"
    distinct_project_names: set = field(default_factory=set)
    incident_sheet_row_count: int = 0
    observation_sheet_row_count: int = 0

    @property
    def all_payloads(self) -> list[RawSafetyEventPayload]:
        return self.incident_payloads + self.observation_payloads

    @property
    def is_structurally_valid(self) -> bool:
        return not self.structure_issues


def _normalize_header(value) -> str:
    return str(value).strip() if value is not None else ""


def _find_sheet(sheet_names: list[str], candidates: tuple[str, ...]) -> str | None:
    by_normalized = {name.strip().lower(): name for name in sheet_names}
    for candidate in candidates:
        match = by_normalized.get(candidate.strip().lower())
        if match is not None:
            return match
    return None


def _header_row(ws) -> list[str]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return [_normalize_header(h) for h in first]


def _header_positions(headers: list[str]) -> dict[str, list[int]]:
    """Every column index a header name appears at -- handles the
    Observation sheet's own duplicate `Status` header correctly (a plain
    `dict(zip(headers, row))` would silently keep only the *last*
    occurrence, discarding the first `Status` column's values entirely)."""
    positions: dict[str, list[int]] = {}
    for i, h in enumerate(headers):
        positions.setdefault(h, []).append(i)
    return positions


def _get(row: tuple, positions: dict[str, list[int]], name: str, occurrence: int = 0):
    idxs = positions.get(name)
    if not idxs or occurrence >= len(idxs) or idxs[occurrence] >= len(row):
        return None
    return row[idxs[occurrence]]


def _clean_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_utc(value) -> datetime | None:
    """Never guess-parses a non-datetime cell -- openpyxl (with
    `data_only=True`) already returns a real `datetime` for a genuine
    Excel date cell; anything else (a stray string, a formula error) is
    reported as missing/invalid by the existing validation pipeline once
    ingested, never fabricated into a plausible-looking date here."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    return None


def _is_blank_row(row: tuple) -> bool:
    return all(c is None for c in row)


def _map_incident_row(headers: list[str], row: tuple) -> RawSafetyEventPayload:
    values = dict(zip(headers, row, strict=False))

    attributes: dict[str, str] = {}
    for column in _INCIDENT_ATTRIBUTE_COLUMNS:
        text = _clean_str(values.get(column))
        if text is not None:
            attributes[column] = text

    return RawSafetyEventPayload(
        event_type=_clean_str(values.get("Incident Type")),
        event_time=_to_utc(values.get("Date")),
        project=_clean_str(values.get("Project")),
        description=_clean_str(values.get("Description")),
        attributes=attributes,
        source_system=SOURCE_SYSTEM,
        source_record_id=_clean_str(values.get("Document No")),
    )


def _map_observation_row(headers: list[str], positions: dict[str, list[int]], row: tuple) -> RawSafetyEventPayload:
    status_primary = _clean_str(_get(row, positions, "Status", occurrence=0))
    status_secondary = _clean_str(_get(row, positions, "Status", occurrence=1))
    observation_channel = _clean_str(_get(row, positions, "Type"))
    recommendation = _clean_str(_get(row, positions, "Recommendation"))
    action_taken = _clean_str(_get(row, positions, "Action Taken"))
    comment = _clean_str(_get(row, positions, "Comment"))

    attributes: dict[str, str] = {}
    if observation_channel is not None:
        attributes["observation_channel"] = observation_channel
    if status_secondary is not None:
        # Never merged with `status_primary` -- see module docstring.
        attributes["status_secondary_raw"] = status_secondary
    if recommendation is not None:
        attributes["recommendation"] = recommendation
    if action_taken is not None:
        attributes["action_taken"] = action_taken
    if comment is not None:
        attributes["comment"] = comment

    return RawSafetyEventPayload(
        event_type="OBSERVATION",
        event_subtype=_clean_str(_get(row, positions, "Category")),
        event_time=_to_utc(_get(row, positions, "Date")),
        status=status_primary,
        project=_clean_str(_get(row, positions, "Project Name")),
        description=_clean_str(_get(row, positions, "Finding")),
        attributes=attributes,
        source_system=SOURCE_SYSTEM,
        source_record_id=_clean_str(_get(row, positions, "Document No")),
    )


def load_real_dataset(path: str | Path) -> RealDatasetLoadResult:
    """Load and map both sheets of the supplied workbook. Never raises
    on a structural problem — every issue (a missing sheet, a missing
    required header, a row that could not be mapped at all) is recorded
    on the returned result instead, so a caller always gets back
    whatever *could* be mapped, exactly like every other SIE ingestion
    entry point already behaves (one bad record never aborts the rest)."""
    result = RealDatasetLoadResult()
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        incident_sheet = _find_sheet(wb.sheetnames, _INCIDENT_SHEET_CANDIDATES)
        observation_sheet = _find_sheet(wb.sheetnames, _OBSERVATION_SHEET_CANDIDATES)

        if incident_sheet is None:
            result.structure_issues.append(
                StructureValidationIssue("<workbook>", "MISSING_INCIDENT_SHEET", "No Incident(s) sheet found.")
            )
        if observation_sheet is None:
            result.structure_issues.append(
                StructureValidationIssue("<workbook>", "MISSING_OBSERVATION_SHEET", "No Observation(s) sheet found.")
            )

        if incident_sheet is not None:
            ws = wb[incident_sheet]
            headers = _header_row(ws)
            missing = [h for h in _INCIDENT_REQUIRED_HEADERS if h not in headers]
            if missing:
                result.structure_issues.append(
                    StructureValidationIssue(
                        incident_sheet, "MISSING_REQUIRED_HEADERS", f"Missing required column(s): {missing}"
                    )
                )
            else:
                for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if row is None or _is_blank_row(row):
                        continue
                    result.incident_sheet_row_count += 1
                    try:
                        payload = _map_incident_row(headers, row)
                    except Exception as exc:  # noqa: BLE001 -- one malformed row must never abort the load
                        result.row_level_skips.append(
                            f"{incident_sheet} row {row_number}: unmappable ({type(exc).__name__})"
                        )
                        continue
                    result.incident_payloads.append(payload)
                    if payload.project:
                        result.distinct_project_names.add(payload.project)

        if observation_sheet is not None:
            ws = wb[observation_sheet]
            headers = _header_row(ws)
            positions = _header_positions(headers)
            missing = [h for h in _OBSERVATION_REQUIRED_HEADERS if h not in positions]
            if missing:
                result.structure_issues.append(
                    StructureValidationIssue(
                        observation_sheet, "MISSING_REQUIRED_HEADERS", f"Missing required column(s): {missing}"
                    )
                )
            else:
                for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if row is None or _is_blank_row(row):
                        continue
                    result.observation_sheet_row_count += 1
                    try:
                        payload = _map_observation_row(headers, positions, row)
                    except Exception as exc:  # noqa: BLE001 -- one malformed row must never abort the load
                        result.row_level_skips.append(
                            f"{observation_sheet} row {row_number}: unmappable ({type(exc).__name__})"
                        )
                        continue
                    result.observation_payloads.append(payload)
                    if payload.project:
                        result.distinct_project_names.add(payload.project)
    finally:
        wb.close()

    return result
