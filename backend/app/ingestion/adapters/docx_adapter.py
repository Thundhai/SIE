"""DOCX adapter — walks the document in true document order (headings,
paragraphs, and tables interleaved as they actually appear), tracking
the most recent heading as the current section so every paragraph/table
extracted carries its `section_title`.

`python-docx` exposes `document.paragraphs` and `document.tables` as two
separate flat lists with no built-in "give me everything in document
order" iterator, so `_iter_block_items` below implements the well-known
recipe for that: walk the underlying XML body's direct children and map
each `<w:p>`/`<w:tbl>` element back to its python-docx wrapper object.
Without this, a table appearing between two paragraphs would be
extracted out of order relative to them, silently corrupting the
"preserve section hierarchy" guarantee this adapter exists to provide.
"""

import io

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.adapters.base import AdapterValidationError, ExtractionResult
from app.ingestion.normalized_content import NormalizedContent, NormalizedContentType


def _iter_block_items(document: Document):
    """Yield each top-level paragraph and table in the document body, in
    the order they actually appear."""
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    if style_name.startswith("Heading"):
        suffix = style_name.replace("Heading", "").strip()
        return int(suffix) if suffix.isdigit() else 1
    if style_name == "Title":
        return 0
    return None


def _list_type(paragraph: Paragraph) -> str | None:
    """Whether this paragraph is a bulleted/numbered list item, detected
    the same way `_heading_level` detects headings: by Word's own named
    paragraph style, not by attempting to parse `numPr` numbering XML
    (see the milestone's "do not attempt perfect structural
    understanding" principle — this covers the common case of Word's
    built-in List Bullet/List Number styles, applied here for section 9
    of the milestone spec: preserving list semantics rather than letting
    each bullet become an arbitrarily separated chunk)."""
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    if style_name.startswith("List Bullet"):
        return "bullet"
    if style_name.startswith("List Number"):
        return "numbered"
    return None


class DOCXAdapter:
    format_id = "docx"
    content_category = "document"

    def detect(self, *, media_type: str, extension: str) -> bool:
        return media_type.endswith("wordprocessingml.document") or extension == ".docx"

    def validate(self, file_bytes: bytes, *, filename: str) -> list[str]:
        try:
            Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise AdapterValidationError(f"Not a valid DOCX document: {exc}") from exc
        return []

    def extract(self, file_bytes: bytes) -> ExtractionResult:
        document = Document(io.BytesIO(file_bytes))
        core_props = document.core_properties

        blocks: list[dict] = []
        current_section: str | None = None
        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                level = _heading_level(item)
                text = item.text.strip()
                if not text:
                    continue
                if level is not None:
                    current_section = text
                    blocks.append({"kind": "heading", "text": text, "level": level})
                else:
                    blocks.append(
                        {
                            "kind": "paragraph",
                            "text": text,
                            "section_title": current_section,
                            "list_type": _list_type(item),
                        }
                    )
            elif isinstance(item, Table):
                rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                blocks.append(
                    {"kind": "table", "rows": rows, "section_title": current_section}
                )

        warnings: list[str] = []
        if not blocks:
            warnings.append("Document contains no extractable paragraphs, headings, or tables.")

        return ExtractionResult(
            raw={"blocks": blocks},
            warnings=warnings,
            metadata={
                "document_title": core_props.title or None,
                "document_author": core_props.author or None,
                "block_count": len(blocks),
            },
        )

    def normalize(self, extraction: ExtractionResult) -> list[NormalizedContent]:
        contents: list[NormalizedContent] = []
        for block in extraction.raw["blocks"]:
            if block["kind"] == "heading":
                contents.append(
                    NormalizedContent(
                        content_type=NormalizedContentType.TEXT,
                        text=block["text"],
                        section_title=block["text"],
                        metadata={"heading_level": block["level"]},
                    )
                )
            elif block["kind"] == "paragraph":
                metadata = {}
                if block["list_type"] is not None:
                    metadata = {"is_list_item": True, "list_type": block["list_type"]}
                contents.append(
                    NormalizedContent(
                        content_type=NormalizedContentType.TEXT,
                        text=block["text"],
                        section_title=block["section_title"],
                        metadata=metadata,
                    )
                )
            elif block["kind"] == "table":
                rows = block["rows"]
                text = "\n".join(" | ".join(row) for row in rows)
                contents.append(
                    NormalizedContent(
                        content_type=NormalizedContentType.TABLE,
                        text=text,
                        section_title=block["section_title"],
                        metadata={"rows": rows, "row_count": len(rows)},
                    )
                )
        return contents
