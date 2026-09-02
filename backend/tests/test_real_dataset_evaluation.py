"""Tests for the Real Enterprise Dataset Evaluation v0.1 pipeline —
`app/validation/real_dataset_loader.py` (workbook -> RawSafetyEventPayload)
and `app/validation/real_dataset_evaluation.py` (the full evaluation
runner: ingestion, data quality, terminology review, temporal integrity,
provenance, intelligence readiness, predictive readiness, the
observation-vs-incident temporal analysis, and HSE review queuing).

Every test here runs against `tests/fixtures/real_dataset_xlsx_fixture.py`'s
**synthetic**, xlsx-shaped workbook — never the real uploaded workbook,
which lives only outside this repository and is never referenced by
path, name, or content from any test (see that fixture module's own
docstring). Passing these tests is not, and must never be described as,
validation against real customer data.

Loader-only tests need no database and run unconditionally. Every test
that runs the full evaluation (ingestion, temporal/provenance/predictive
checks) requires real PostgreSQL (`@requires_postgres`, `pg_session`) —
same rationale as `tests/test_enterprise_dataset_validation.py`.
"""

from __future__ import annotations

import json
from datetime import datetime, timezone

from app.intelligence.features import compute_feature_set
from app.intelligence.temporal import events_as_of
from app.intelligence.terminology_review import TerminologyReviewStatus
from app.validation.real_dataset_evaluation import (
    run_real_dataset_evaluation,
    to_dict,
    to_markdown,
)
from app.validation.real_dataset_loader import load_real_dataset
from tests.fixtures.real_dataset_xlsx_fixture import (
    OBSERVATION_ONLY_PROJECT,
    SYNTHETIC_PII_NAME,
    build_synthetic_workbook,
)
from tests.postgres_support import requires_postgres

_FIXED_BASE_TIME = datetime(2026, 6, 1, tzinfo=timezone.utc)


# --- Loader-only tests (no database) --------------------------------------------------------


def test_workbook_structure_validation_reports_missing_sheet(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "no_incident_sheet.xlsx"
    wb = Workbook()
    wb.active.title = "Observation"
    wb.active.append(["Document No", "Date", "Project Name", "Category"])
    wb.save(str(path))

    result = load_real_dataset(path)
    codes = {issue.code for issue in result.structure_issues}
    assert "MISSING_INCIDENT_SHEET" in codes
    assert result.is_structurally_valid is False


def test_workbook_structure_validation_reports_missing_required_header(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "missing_header.xlsx"
    wb = Workbook()
    wb.active.title = "Incident"
    wb.active.append(["Document No", "Date"])  # missing Project, Incident Type
    wb.create_sheet("Observation").append(["Document No", "Date", "Project Name", "Category"])
    wb.save(str(path))

    result = load_real_dataset(path)
    codes = {issue.code for issue in result.structure_issues}
    assert "MISSING_REQUIRED_HEADERS" in codes


def test_a_well_formed_synthetic_workbook_has_no_structure_issues(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)
    assert result.structure_issues == []
    assert result.row_level_skips == []


def test_incident_mapping_preserves_unresolved_terminology_and_provenance(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    chars = build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    assert result.incident_sheet_row_count == chars.incident_data_row_count
    by_id = {p.source_record_id: p for p in result.incident_payloads if p.source_record_id}
    mapped_row = by_id["SYN-IC-0001"]
    assert mapped_row.event_type == chars.mapped_incident_type_term  # "Injury" -- never translated by the loader
    unknown_row = by_id["SYN-IC-0002"]
    assert unknown_row.event_type == chars.unknown_incident_type_term  # "NearMiss" -- passed through unresolved
    assert mapped_row.source_system == "alm-hse-xlsx"
    assert mapped_row.project in chars.incident_projects


def test_observation_mapping_never_merges_the_duplicate_status_columns(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    by_id = {p.source_record_id: p for p in result.observation_payloads if p.source_record_id}
    closed_row = by_id["SYN-OB-0001"]
    assert closed_row.status == "Closed"
    assert closed_row.attributes["status_secondary_raw"] == "REVIEW REQUIRED"
    open_row = by_id["SYN-OB-0002"]
    assert open_row.status == "Open"
    assert open_row.attributes["status_secondary_raw"] == "OK"
    # Never silently coerced to match -- primary and secondary can and do disagree in shape.
    assert closed_row.status != closed_row.attributes["status_secondary_raw"]


def test_missing_document_no_and_missing_date_never_drop_a_row(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    chars = build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    missing_id_incidents = [p for p in result.incident_payloads if p.source_record_id is None]
    missing_time_incidents = [p for p in result.incident_payloads if p.event_time is None]
    assert len(missing_id_incidents) == chars.missing_document_no_incident_rows
    assert len(missing_time_incidents) == chars.missing_date_incident_rows

    missing_id_observations = [p for p in result.observation_payloads if p.source_record_id is None]
    missing_time_observations = [p for p in result.observation_payloads if p.event_time is None]
    assert len(missing_id_observations) == chars.missing_document_no_observation_rows
    assert len(missing_time_observations) == chars.missing_date_observation_rows


def test_duplicate_document_no_produces_two_distinct_payloads_at_the_loader_stage(tmp_path):
    """The loader never deduplicates -- that decision belongs entirely to
    the existing, unmodified ingestion pipeline (SKIPPED_IDEMPOTENT /
    version handling), never guessed here."""
    path = tmp_path / "synthetic.xlsx"
    chars = build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    dup_incidents = [p for p in result.incident_payloads if p.source_record_id == chars.duplicate_incident_document_no]
    dup_observations = [
        p for p in result.observation_payloads if p.source_record_id == chars.duplicate_observation_document_no
    ]
    assert len(dup_incidents) == 2
    assert len(dup_observations) == 2


def test_future_timestamp_row_is_loaded_as_is_never_filtered(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    future_rows = [p for p in result.incident_payloads if p.event_time and p.event_time > _FIXED_BASE_TIME]
    assert len(future_rows) >= 1


def test_project_reconciliation_the_incident_project_set_is_a_subset_of_the_observation_project_set(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    chars = build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    incident_projects = {p.project for p in result.incident_payloads if p.project}
    observation_projects = {p.project for p in result.observation_payloads if p.project}
    assert incident_projects == chars.incident_projects
    assert observation_projects == chars.observation_projects
    assert incident_projects <= observation_projects  # every incident-bearing project also has observations
    assert OBSERVATION_ONLY_PROJECT in (observation_projects - incident_projects)
    assert result.distinct_project_names == incident_projects | observation_projects


def test_pii_bearing_narrative_fields_are_preserved_on_the_payload_never_stripped(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    chars = build_synthetic_workbook(path, base_time=_FIXED_BASE_TIME)
    result = load_real_dataset(path)

    pii_incident = next(p for p in result.incident_payloads if p.source_record_id == chars.pii_bearing_incident_document_no)
    assert SYNTHETIC_PII_NAME in (pii_incident.description or "")
    assert SYNTHETIC_PII_NAME in pii_incident.attributes.get("PeopleInvolved", "")

    pii_observation = next(
        p for p in result.observation_payloads if p.source_record_id == chars.pii_bearing_observation_document_no
    )
    assert SYNTHETIC_PII_NAME in (pii_observation.description or "")


# --- Full evaluation runner tests (real PostgreSQL) -----------------------------------------


def _run(pg_session, tmp_path, org_name: str, *, base_time=_FIXED_BASE_TIME, as_of=None):
    path = tmp_path / f"{org_name.replace(' ', '_')}.xlsx"
    build_synthetic_workbook(path, base_time=base_time)
    return run_real_dataset_evaluation(
        pg_session, workbook_path=path, organization_name=org_name,
        label="SYNTHETIC -- tests/fixtures/real_dataset_xlsx_fixture.py, not real customer data.",
        as_of=as_of,
    )


@requires_postgres
def test_full_evaluation_ingests_every_mappable_row_and_creates_one_site_per_project(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - Full Ingestion")
    assert report.workbook_structure_issues == []
    assert report.row_level_skips == []
    assert len(report.project_to_site_id) == 3  # Alpha, Beta, Gamma
    assert report.validation.observed.total_records == report.incident_sheet_row_count + report.observation_sheet_row_count


@requires_postgres
def test_terminology_review_flags_the_unknown_incident_term_and_never_guesses_it(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - Terminology")
    entries = report.validation.terminology_entries
    unknown = [e for e in entries if e.status == TerminologyReviewStatus.UNKNOWN and e.source_term == "NearMiss"]
    assert len(unknown) == 1
    assert unknown[0].requires_review is True
    assert unknown[0].proposed_canonical_term is None  # never guessed

    mapped = [e for e in entries if e.status == TerminologyReviewStatus.MAPPED and e.source_term == "Injury"]
    assert len(mapped) == 1


@requires_postgres
def test_temporal_integrity_detects_the_future_dated_row_and_passes_the_leakage_check(pg_session, tmp_path):
    # `future_timestamps` is counted against real wall-clock "now"
    # (`app/validation/enterprise_dataset_validation.py`'s own
    # `_count_future_timestamps()`), not against the fixture's own
    # `base_time` -- so this one test builds its workbook relative to
    # real "now" (the default `build_synthetic_workbook()` uses when no
    # `base_time` is given) rather than `_FIXED_BASE_TIME`, so its
    # +30-day row is genuinely in the future.
    report = _run(pg_session, tmp_path, "Real Eval Test - Temporal Integrity", base_time=None)
    assert report.validation.data_quality.future_timestamps >= 1
    assert report.validation.temporal.as_of_leakage_check_passed is True


@requires_postgres
def test_a_genuinely_earlier_as_of_correctly_shows_reduced_visibility_not_a_leak(pg_session, tmp_path):
    """The literal "late-arriving historical event must not affect
    earlier analysis" regression test, run through the real evaluation
    entry point (`run_real_dataset_evaluation`), not just the underlying
    harness: every row in this workbook is ingested "now" (real wall
    clock), so a caller-supplied `as_of` from long before that moment
    must see none of it -- `ingestion_time <= as_of` failing for every
    row is the correct behavior, not a bug."""
    ancient_as_of = datetime(2000, 1, 1, tzinfo=timezone.utc)
    report = _run(pg_session, tmp_path, "Real Eval Test - Ancient As-of", as_of=ancient_as_of)
    assert report.validation.intelligence_readiness.event_count == 0
    assert report.validation.intelligence_readiness.data_sufficiency == "INSUFFICIENT_DATA"
    assert report.validation.predictive_readiness.eligible_sites == 0


@requires_postgres
def test_provenance_chains_are_intact_end_to_end(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - Provenance")
    pv = report.validation.provenance
    # checked_count excludes only the two REJECTED rows (missing
    # Document No, on each sheet -- no source_record_id means no
    # canonical event to check provenance against at all). The two
    # missing-Date rows still get a canonical event (QUARANTINED, not
    # REJECTED) and so are still checked -- an existing pipeline
    # distinction this test asserts on, not assumes.
    assert pv.checked_count == report.validation.observed.total_records - report.validation.data_quality.missing_identifiers
    assert pv.checked_count > 0
    assert pv.all_chains_intact is True
    assert pv.failures == []


@requires_postgres
def test_observation_incident_temporal_analysis_totals_match_the_datasets_own_valid_counts(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - Obs-Incident Temporal")
    oit = report.observation_incident_temporal
    assert "RETROSPECTIVE" in oit.methodology_note
    assert "causal" in oit.methodology_note.lower()

    total_observations = sum(site.total_observations for site in oit.per_site)
    total_incidents = sum(site.total_incidents for site in oit.per_site)
    # Every VALID event_time in this fixture falls within the analysis
    # window (base_time is recent, the window covers 240 days back), so
    # the retrospective totals are exactly the dataset's own VALID,
    # distinct-canonical-event counts: of 7 observation rows, 1
    # (missing Document No) is REJECTED (no canonical event at all) and
    # 1 (missing Date) is QUARANTINED (excluded from this
    # non-quarantined analysis); the exact-duplicate Document No pair
    # (SYN-OB-0005 appearing twice, no version field) upserts onto ONE
    # canonical event, not two -- leaving exactly 4: SYN-OB-0001,
    # SYN-OB-0002, the collapsed SYN-OB-0005, SYN-OB-0006. The same
    # shape holds for incidents: 8 rows minus 1 missing-Date
    # (QUARANTINED) minus 1 unknown-terminology "NearMiss" (QUARANTINED)
    # minus the SYN-IC-0005 duplicate pair collapsing to one, leaving
    # exactly 4: SYN-IC-0001, the collapsed SYN-IC-0005, SYN-IC-0006,
    # SYN-IC-0007 (the missing-Document-No row is REJECTED).
    assert total_observations == 4
    assert total_incidents == 4

    # Gamma is observation-only -- must show zero incidents, never a guessed nonzero.
    gamma = next(s for s in oit.per_site if s.project_name == OBSERVATION_ONLY_PROJECT)
    assert gamma.total_incidents == 0
    assert gamma.total_observations > 0


@requires_postgres
def test_predictive_readiness_never_crashes_and_reports_a_coherent_shape(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - Predictive Readiness")
    pr = report.validation.predictive_readiness
    assert pr.eligible_sites + pr.excluded_sites == len(report.project_to_site_id)
    # A live "as_of=now" evaluation of just-ingested historical data can
    # never have a confirmed label yet -- the 30-day horizon has not
    # elapsed as of real "now" for any of this fixture's dates.
    assert pr.positive_labels == 0
    assert pr.negative_labels == 0


@requires_postgres
def test_pii_is_never_read_by_feature_engineering(pg_session, tmp_path):
    """`compute_feature_set()` only ever reads structured fields
    (event_type/severity/site/time) -- this asserts that black-box, by
    ingesting a fixture whose narrative fields carry a fabricated name
    and checking the computed FeatureValue set never surfaces it."""
    report = _run(pg_session, tmp_path, "Real Eval Test - PII Features")
    site_id = next(iter(report.project_to_site_id.values()))
    import uuid as _uuid

    query = events_as_of(organization_id=report.organization_id, as_of=report.generated_at, site_id=_uuid.UUID(site_id))
    events = pg_session.execute(query).scalars().all()
    assert events  # sanity: the site actually has events to compute over

    features = compute_feature_set(events, window_days=365, as_of=report.generated_at)
    serialized = json.dumps({k: v.value for k, v in features.items()}, default=str)
    assert SYNTHETIC_PII_NAME not in serialized
    # Structural guarantee, not just an absence-of-evidence check: no
    # FeatureValue field is ever a free-text string in the first place.
    for fv in features.values():
        assert fv.value is None or isinstance(fv.value, (int, float))


@requires_postgres
def test_pii_never_appears_in_the_rendered_machine_or_human_readable_report(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - PII Report")
    serialized = json.dumps(to_dict(report), default=str) + "\n" + to_markdown(report)
    assert SYNTHETIC_PII_NAME not in serialized


@requires_postgres
def test_tenant_isolation_between_two_independently_run_evaluations(pg_session, tmp_path):
    report_a = _run(pg_session, tmp_path, "Real Eval Test - Tenant A")
    report_b = _run(pg_session, tmp_path, "Real Eval Test - Tenant B")

    assert report_a.organization_id != report_b.organization_id
    assert set(report_a.project_to_site_id.values()).isdisjoint(report_b.project_to_site_id.values())
    assert report_a.validation.provenance.checked_count == report_b.validation.provenance.checked_count
    # Org B's own predictive readiness must never reflect org A's sites.
    assert report_b.validation.predictive_readiness.eligible_sites + report_b.validation.predictive_readiness.excluded_sites == 3


@requires_postgres
def test_intelligence_reproducibility_across_two_independent_runs_of_the_same_fixture(pg_session, tmp_path):
    """Two organizations, each freshly evaluating its own independently
    built copy of the identical deterministic synthetic workbook, must
    yield identical ingestion/data-quality/terminology summaries -- the
    evaluation runner itself introduces no run-to-run nondeterminism."""
    report_1 = _run(pg_session, tmp_path, "Real Eval Test - Repro 1")
    report_2 = _run(pg_session, tmp_path, "Real Eval Test - Repro 2")

    assert report_1.validation.ingestion == report_2.validation.ingestion
    assert report_1.validation.data_quality == report_2.validation.data_quality
    assert report_1.validation.terminology_summary == report_2.validation.terminology_summary
    assert report_1.incident_sheet_row_count == report_2.incident_sheet_row_count
    assert report_1.observation_sheet_row_count == report_2.observation_sheet_row_count


@requires_postgres
def test_review_candidates_are_queued_for_the_unknown_terminology_and_quarantined_records(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - HSE Review Queue")
    assert report.queued_review_candidate_count > 0


@requires_postgres
def test_report_never_claims_predictive_accuracy(pg_session, tmp_path):
    report = _run(pg_session, tmp_path, "Real Eval Test - No Accuracy Claim")
    md = to_markdown(report)
    # The report explicitly *disclaims* PR-AUC/ROC-AUC/precision/recall
    # (app/validation/report.py's own predictive-readiness section) --
    # mentioning those terms in that disclaimer is correct and expected.
    # What must never appear is an actual reported *value* for one.
    assert "No predictive accuracy, PR-AUC, ROC-AUC, precision, recall" in md
    md_lower = md.lower()
    for forbidden in ("pr-auc:", "roc-auc:", "precision:", "recall:", "accuracy achieved", "accuracy:"):
        assert forbidden not in md_lower
